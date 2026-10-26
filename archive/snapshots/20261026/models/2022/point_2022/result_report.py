"""Readable outputs and auditable residual diagnostics for the point model."""
from pathlib import Path
from datetime import date, datetime, timedelta
from collections import defaultdict
import csv
import json
import math
import statistics


def read_table(path):
    with path.open(encoding='utf-8-sig', newline='') as stream:
        return list(csv.DictReader(stream))


def included(row):
    return row['included_in_metrics'] in (True, 'True')


def diagnose(archive):
    pairs = read_table(archive/'observation_pairs.csv')
    events = read_table(archive/'event_parameters.csv')
    config = json.loads((archive/'config_snapshot.json').read_text(encoding='utf-8'))
    details = []
    summaries = []
    for treatment in ('RM', 'RCC'):
        groups = defaultdict(list)
        for row in pairs:
            if row['treatment'] == treatment and included(row):
                groups[row['date']].append(row)
        errors = []
        for day, rows in groups.items():
            observed = statistics.mean(float(r['value']) for r in rows)
            predicted = float(rows[0]['simulated_TN_mg_L'])
            errors.append(dict(treatment=treatment, date=day, observed=observed,
                               predicted=predicted, residual=predicted-observed,
                               squared_error=(predicted-observed)**2))
        total = sum(r['squared_error'] for r in errors)
        errors.sort(key=lambda r: r['squared_error'], reverse=True)
        for rank, row in enumerate(errors, 1):
            row.update(rank=rank, squared_error_share_pct=100*row['squared_error']/total if total else 0.)
        details.extend(errors)
        late = [r['observed'] for r in errors if r['date'] >= '2022-08-01']
        summaries.append(dict(treatment=treatment, scored_dates=len(errors),
                              top3_error_share_pct=sum(r['squared_error_share_pct'] for r in errors[:3]),
                              late_observed_median=statistics.median(late) if late else None,
                              worst=errors[0] if errors else None))
    components = []
    for event in events:
        coefficients = config['coefficients'][event['stage']]
        nitrogen = coefficients['aN']*float(event['N_kg_ha'])
        soil = coefficients['apH']*float(event['pH']) + coefficients['aSOM']*float(event['SOM_percent_assumed'])
        components.append(dict(treatment=event['treatment'], date=event['date'], stage=event['stage'],
                               N_term_mg_L=nitrogen, constant_term_mg_L=coefficients['a0'],
                               soil_term_mg_L=soil, non_N_term_mg_L=coefficients['a0']+soil,
                               predicted_C0_mg_L=float(event['C0_mg_L'])))
    # Endpoint decay checks describe the observed first-event decline, not fitted parameters.
    decay_checks = []
    for treatment in ('RM', 'RCC'):
        event = min((e for e in events if e['treatment']==treatment), key=lambda e:e['date'])
        first_day = date.fromisoformat(event['date'])
        points = sorted([r for r in details if r['treatment']==treatment and
                         1 <= (date.fromisoformat(r['date'])-first_day).days <= 7], key=lambda r:r['date'])
        if len(points)>=2 and points[0]['observed']>0 and points[-1]['observed']>0:
            elapsed = (date.fromisoformat(points[-1]['date'])-date.fromisoformat(points[0]['date'])).days
            decay_checks.append(dict(treatment=treatment, start=points[0]['date'], end=points[-1]['date'],
                                     apparent_k=math.log(points[0]['observed']/points[-1]['observed'])/elapsed,
                                     model_k=float(event['k_day'])))
    result = dict(summary=summaries, error_contributions=details,
                  C0_components=components, endpoint_decay_checks=decay_checks,
                  interpretation='diagnostics on scored date means; parameters remain unchanged')
    (archive/'diagnostics.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return result


def make_plot(destination, daily, pairs, events, scores, calibrated=False):
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.font_manager import FontProperties
    import matplotlib.dates as mdates
    from matplotlib import rc_context
    font = FontProperties(fname='C:/Windows/Fonts/msyh.ttc').get_name()
    observation_dates = [date.fromisoformat(r['date']) for r in pairs]
    # Focus on the monitoring interval; the full simulation remains in the workbook.
    start = min(observation_dates)-timedelta(days=1)
    end = max(observation_dates)+timedelta(days=2)
    last_event = max(date.fromisoformat(e['date']) for e in events)
    split = min(end, max(start+timedelta(days=1), last_event+timedelta(days=7)))
    windows = [(start,split), (split,end)] if split<end else [(start,end)]
    with rc_context({'font.family':font,'axes.unicode_minus':False,'font.size':10}):
        fig = Figure(figsize=(12,7), constrained_layout=True)
        FigureCanvasAgg(fig)
        axes = fig.subplots(2,len(windows),squeeze=False)
        score_map = {r['treatment']:r for r in scores}
        for row_index,treatment in enumerate(('RM','RCC')):
            records = [r for r in daily if r['treatment']==treatment]
            obs = [r for r in pairs if r['treatment']==treatment]
            for col,(left,right) in enumerate(windows):
                ax = axes[row_index,col]
                selected = [r for r in records if left<=date.fromisoformat(r['date'])<=right]
                ax.plot([date.fromisoformat(r['date']) for r in selected], [float(r['TN_mg_L']) for r in selected],
                        color='#256387', lw=1.6, label='栅格驱动预测')
                for scored,marker,label in [(True,'o','参与评价的观测'),(False,'x','待核或未计入评价的观测')]:
                    selected_obs = [r for r in obs if included(r)==scored and left<=date.fromisoformat(r['date'])<=right]
                    ax.scatter([date.fromisoformat(r['date']) for r in selected_obs], [float(r['value']) for r in selected_obs],
                               marker=marker, s=23, color='#222222' if scored else '#c78b37', label=label, zorder=3)
                for event in events:
                    event_day = date.fromisoformat(event['date'])
                    if event['treatment']==treatment and left<=event_day<=right:
                        ax.axvline(event_day,color='#929292',ls=':',lw=.8)
                        ax.text(event_day,.96,'施肥',transform=ax.get_xaxis_transform(),ha='center',va='top',fontsize=8)
                score = score_map[treatment]
                title = ('RM 水稻单作' if treatment=='RM' else 'RCC 稻虾共作')
                title += '｜施肥高值期' if col==0 else '｜后期低值期（独立纵轴）'
                ax.set_title(title,loc='left',fontsize=11)
                ax.set_ylabel('TN浓度（mg/L）')
                ax.set_xlim(left,right)
                ax.set_ylim(bottom=0)
                ax.xaxis.set_major_locator(mdates.DayLocator(interval=7 if col==0 else 14))
                ax.xaxis.set_major_formatter(mdates.DateFormatter('%m-%d'))
                ax.grid(axis='y',alpha=.2)
                for side in ('top','right'):
                    ax.spines[side].set_visible(False)
                if col==0:
                    rmse = f'{float(score["RMSE_mg_L"]):.2f}' if score['RMSE_mg_L'] else '—'
                    nse = f'{float(score["NSE"]):.2f}' if score['NSE'] else '—'
                    ax.text(.98,.84,f'RMSE={rmse} mg/L\nNSE={nse}；n={int(score["n"])}',
                            transform=ax.transAxes,ha='right',va='top',fontsize=9)
                if row_index==0 and col==len(windows)-1:
                    ax.legend(loc='upper right',fontsize=8,frameon=False)
        fig.suptitle('2022年点位TN浓度：栅格C0与区间率定k' if calibrated else '2022年点位TN浓度：栅格驱动预测与观测',fontsize=15)
        fig.supxlabel('月—日；指标含率定日期；右侧纵轴单独缩放；留出检验见Excel' if calibrated else '月—日；指标基于入选日期，右侧纵轴单独缩放；完整6—10月模拟见Excel',fontsize=9)
        fig.savefig(destination,dpi=180)


def export_workbook(archive, destination, diagnostics):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    scores = read_table(archive/'metrics.csv')
    pairs = read_table(archive/'observation_pairs.csv')
    events = read_table(archive/'event_parameters.csv')
    daily = read_table(archive/'daily_simulation.csv')
    nitrogen = read_table(archive/'nitrogen_input.csv')[0]
    workbook = Workbook()
    workbook.remove(workbook.active)

    def table(title, rows):
        sheet = workbook.create_sheet(title)
        if not rows:
            return
        sheet.append(list(rows[0]))
        for row in rows:
            sheet.append(list(row.values()))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int,float)) and not isinstance(cell.value,bool):
                    cell.number_format='0.00'
                elif isinstance(cell.value, str):
                    cell.data_type='s'
        # Ordinary data sheets, with enough width to read field names and values.
        for column in sheet.columns:
            width = max(sum(2 if ord(c)>255 else 1 for c in str(cell.value or '')) for cell in column)
            sheet.column_dimensions[column[0].column_letter].width=min(70,max(12,width+2))
        if title in ('误差诊断', '复现位置'):
            for row in sheet.iter_rows(min_row=2):
                lines = 1
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
                    width = sheet.column_dimensions[cell.column_letter].width-2
                    text_width = sum(2 if ord(c)>255 else 1 for c in str(cell.value or ''))
                    lines = max(lines, math.ceil(text_width/width))
                sheet.row_dimensions[row[0].row].height=15*lines

    table('验证指标',[{'处理':r['treatment'],'有效日期数':int(r['n']), 'RMSE_mg_L':float(r['RMSE_mg_L']) if r['RMSE_mg_L'] else None,
                         'MAE_mg_L':float(r['MAE_mg_L']) if r['MAE_mg_L'] else None,
                         '平均偏差_mg_L':float(r['bias_mg_L']) if r['bias_mg_L'] else None,
                         'NSE':float(r['NSE']) if r['NSE'] else None,
                         'Pearson相关系数平方':float(r['Pearson_r2']) if r['Pearson_r2'] else None} for r in scores])
    reasons = {'included':'参与评价','date_pending':'读图日期待核','value_pending':'观测值待核',
               'fertilizer_sampling_order_pending':'施肥与采样先后待核','outside_simulation':'模拟时段外','initial_TN_pending':'初始浓度待核'}
    table('观测对照',[{'处理':r['treatment'],'日期':r['date'],'观测TN_mg_L':float(r['value']),
                       '模拟TN_mg_L':float(r['simulated_TN_mg_L']) if r['simulated_TN_mg_L'] else None,
                       '模拟减观测_mg_L':float(r['residual_mg_L']) if r['residual_mg_L'] else None,
                       '参与评价':'是' if included(r) else '否',
                       '数据用途':r.get('calibration_role','文献参数预测'),
                       '说明':'；'.join(reasons.get(x,x) for x in r['evaluation_status'].split(';')),
                       '日期下界':r['date_lower'],'日期上界':r['date_upper'],'记录编号':r['record_id']} for r in pairs])
    by_date = defaultdict(dict)
    for row in daily:
        by_date[row['date']][row['treatment']]=float(row['TN_mg_L'])
    table('逐日模拟',[{'日期':day,'RM模拟TN_mg_L':v.get('RM'),'RCC模拟TN_mg_L':v.get('RCC')} for day,v in sorted(by_date.items())])
    component_map = {(r['treatment'],r['date']):r for r in diagnostics['C0_components']}
    table('施肥与浓度',[{'处理':r['treatment'],'施肥日期':r['date'],'阶段':{'basal':'基肥','tillering':'蘖肥','panicle':'穗肥'}[r['stage']],
                        '施肥比例':float(r['N_fraction']),'事件施氮代理量_kgN_ha':float(r['N_kg_ha']),
                        '初始浓度C0_mg_L':float(r['C0_mg_L']),'施氮项_mg_L':component_map[r['treatment'],r['date']]['N_term_mg_L'],
                        '常数及土壤项_mg_L':component_map[r['treatment'],r['date']]['non_N_term_mg_L'],
                        '衰减系数_每天':float(r['k_day']),'响应日期':r['response_date']} for r in events])
    diagnosis_rows = []
    for summary in diagnostics['summary']:
        worst = summary['worst']
        if worst:
            diagnosis_rows.append({'项目':summary['treatment']+'峰值偏差','证据':f"{worst['date']}：预测{worst['predicted']:.2f}，观测{worst['observed']:.2f} mg/L；最大3个日期贡献{summary['top3_error_share_pct']:.2f}%平方误差。",
                                   '判断':'当前误差主要集中在高值阶段。','下一步':'优先核对C0回归在该点位的适用性及输入空间口径。'})
    for check in diagnostics['endpoint_decay_checks']:
        diagnosis_rows.append({'项目':check['treatment']+'早期衰减','证据':f"{check['start']}至{check['end']}端点表观k={check['apparent_k']:.2f}/天；模型k={check['model_k']:.2f}/天。",
                               '判断':'端点表观衰减与模型k分别列示，受起始幅度影响。','下一步':'结合C0偏差和留出事件判断参数适用性。'})
    diagnosis_rows.extend([
        {'项目':'栅格空间与时间口径','证据':f"2018年度土地施氮{float(nitrogen['annual_N_kg_ha_land']):.2f} kg N/ha土地/年，稻季映射系数{float(nitrogen['rice_season_factor']):.2f}，用于2022年点位。",'判断':'是跨年份、土地到田块的代理输入。','下一步':'继续保留栅格驱动，建立水稻季及面积转换依据。'},
        {'项目':'土壤及回归截距','证据':'蘖肥公式的常数和土壤项见“施肥与浓度”；SOM单位仍按百分数假设。','判断':'峰值误差不能全部归因于施氮栅格。','下一步':'核实SOM来源、土壤代表性，再评估局地C0参数。'},
        {'项目':'日期与阶段','证据':'各事件施肥日和C0响应日见“施肥与浓度”。RM第二次施肥按用户确认的正文日期7月6日；追肥映射蘖肥。','判断':'7月7日峰值记录仍按源表标记为插值日期。','下一步':'核对采样日期依据及追肥阶段归属。'},
        {'项目':'后期背景浓度','证据':'; '.join(f"{s['treatment']}：8月起入选观测中位数{s['late_observed_median']:.2f} mg/L" for s in diagnostics['summary'] if s['late_observed_median'] is not None),
         '判断':'事件指数衰减趋近0，无法表达持续背景浓度。','下一步':'背景项需有独立依据或单独校准，并与验证样本分开。'},
        {'项目':'水量耦合','证据':'当前降雨、T已提取，但不进入TN经验浓度公式。','判断':'现阶段不能逐日表达稀释、换水及田沟混合。','下一步':'浓度输入核定后，再衔接水量过程。'},
    ])
    table('误差诊断',diagnosis_rows)
    calibration_path=archive/'k_calibration.json'
    if calibration_path.exists():
        calibration=json.loads(calibration_path.read_text(encoding='utf-8'))
        table('参数率定',[{'阶段':r['stage'],'文献k':r['literature_k'],'率定k':r['calibrated_k'],
                          '文献下限':r['lower'],'文献上限':r['upper'],'率定日期数':r['n_dates'],
                          '达到边界':'是' if r['at_bound'] else '否',
                          '说明':'固定栅格C0，以首次施肥后第1—7天日期明确的观测最小化浓度平方误差。边界最优不代表峰值已解决。'} for r in calibration['parameters']])
        table('率定与留出',[{'处理':r['treatment'],'评价范围':r['scope'],'参数方案':r['parameter_set'],'日期数':int(r['n']),
                           'RMSE_mg_L':float(r['RMSE_mg_L']) if r['RMSE_mg_L'] else None,
                           'NSE':float(r['NSE']) if r['NSE'] else None,
                           '说明':'同一点事件留出；日期数不足时不能判断迁移精度。'} for r in read_table(archive/'calibration_comparison.csv')])
    table('复现位置',[{'项目':'完整运行记录','内容':str(archive)},
                     {'项目':'参数快照','内容':str(archive/'config_snapshot.json')},
                     {'项目':'模型范围','内容':'田面水TN浓度的栅格驱动迁移检查；观测为读图数据。'},
                     {'项目':'数据精度','内容':'Excel显示两位小数；计算与CSV保留完整精度。'},
                     {'项目':'图件范围','内容':'显示监测时段；完整6—10月模拟保留在逐日模拟表。'}])
    workbook.save(destination)


def publish_results(archive, public_output):
    archive, public_output = Path(archive), Path(public_output)
    public_output.mkdir(parents=True, exist_ok=True)
    diagnostics = diagnose(archive)
    make_plot(public_output/'TN对照图.png',read_table(archive/'daily_simulation.csv'),
              read_table(archive/'observation_pairs.csv'),read_table(archive/'event_parameters.csv'),read_table(archive/'metrics.csv'),
              calibrated=(archive/'k_calibration.json').exists())
    export_workbook(archive,public_output/'2022点位验证结果.xlsx',diagnostics)
    return diagnostics
