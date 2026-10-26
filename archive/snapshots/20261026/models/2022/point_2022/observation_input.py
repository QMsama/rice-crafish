"""将指定田面水TN观测表读入模型，保留日期依据和工作簿行号。"""
from datetime import date, datetime
from hashlib import sha256
import math


def load_workbook_observations(path, sheet_name):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        required = ('日期', '处理代码', '处理', 'TN (mg/L)', '日期依据')
        if not all(name in headers for name in required):
            raise ValueError(f'{sheet_name}需要列：{required}')
        digest = sha256(path.read_bytes()).hexdigest()
        records = []
        statuses = {'刻度': 'printed_tick', '插值': 'between_labels',
                    '外推': 'outside_label_range'}
        for row_number, values in enumerate(rows, 2):
            if all(value is None for value in values):
                continue
            row = dict(zip(headers, values))
            raw_day = row['日期']
            if isinstance(raw_day, datetime):
                day = raw_day.date().isoformat()
            elif isinstance(raw_day, date):
                day = raw_day.isoformat()
            else:
                day = date.fromisoformat(str(raw_day).strip()).isoformat()
            treatment = str(row['处理代码']).strip()
            if treatment not in ('RM', 'RCC'):
                raise ValueError(f'第{row_number}行处理代码不支持：{treatment}')
            value = float(row['TN (mg/L)'])
            if not math.isfinite(value) or value < 0:
                raise ValueError(f'第{row_number}行TN值无效')
            note = str(row['日期依据']).strip()
            if note not in statuses:
                raise ValueError(f'第{row_number}行日期依据需要核查：{note}')
            exact = note == '刻度'
            records.append(dict(
                record_id=f'TN_sheet_row_{row_number:03d}', treatment=treatment,
                metric='TN', date=day, value=value, unit='mg/L',
                date_status=statuses[note], value_status='user_revised_workbook',
                date_exact=day if exact else '', date_lower=day if exact else '',
                date_upper=day if exact else '', role='observed',
                source_family='user_revised_digitized_observations',
                source_file=str(path), source_sheet=sheet_name, source_row=row_number,
                source_sha256=digest, source_date_note=note))
        if not records:
            raise ValueError('指定观测表为空')
        # 插值日期保持待核；最近刻度仅提供日期括区，不代表新增采样记录。
        for record in records:
            if record['date_status'] == 'printed_tick':
                continue
            ticks = sorted({r['date'] for r in records if
                            r['treatment'] == record['treatment'] and
                            r['date_status'] == 'printed_tick'})
            before = [d for d in ticks if d < record['date']]
            after = [d for d in ticks if d > record['date']]
            record['date_lower'] = max(before) if before else ''
            record['date_upper'] = min(after) if after else ''
        return records
    finally:
        workbook.close()
