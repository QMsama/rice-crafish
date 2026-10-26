"""在已归档施肥峰值模型上增加背景及投饵经验项，生成独立情景版本。"""
from pathlib import Path
from datetime import date, datetime
from collections import defaultdict
import csv
import json
import math
import hashlib
import os
import sys
import shutil

sys.dont_write_bytecode=True
HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
os.environ['MPLCONFIGDIR']=str(HERE/'cache/matplotlib')

# 用户参数区：输入固定为已审阅的施肥峰值校准版本。
INPUT_ARCHIVE=HERE/'joint_runs/20261008_221451_815334'
OUTPUT_DIRECTORY=ROOT/'模型/2026.9.3/2022点位验证'
BACKGROUND_FIT_START='2022-07-20'  # 距RM末次施肥已有14天
BACKGROUND_FIT_END='2022-07-31'
HOLDOUT_START='2022-08-01'
FEED_KG_PER_DAY=3.25             # 情景值：论文3.0—3.5 kg/天的中值
FEED_CRUDE_PROTEIN_FRACTION=0.32 # 以投喂质量相同基准解释；待饲料标签核实
PROTEIN_TO_N_FACTOR=6.25
RCC_SYSTEM_AREA_HA=0.31          # 全系统面积，含虾沟；不将全部饲料氮视为进入田面水
FEED_START='2022-06-01'          # 情景覆盖模拟期，不是已核实的实际开喂日期
FEED_END='2022-10-31'            # 情景覆盖模拟期，不是已核实的实际停喂日期
FEED_DAILY_CSV=None             # 可指定CSV，列为date、feed_kg；须覆盖全部模拟日期，停喂填0
# RCC基础背景先与RM共享：固定这一结构假设，才能单独估计投饵附加项。
# 采用日尺度经验项gamma*feed_N_kg_ha_day，不模拟饲料氮库、吸收和水量交换。


def read(path):
    with path.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))


def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()


def nitrogen_from_feed(feed_kg,protein=FEED_CRUDE_PROTEIN_FRACTION,area=RCC_SYSTEM_AREA_HA):
    if not math.isfinite(feed_kg) or feed_kg<0 or not 0<protein<=1 or area<=0:
        raise ValueError('投饵量、蛋白比例或面积无效')
    return feed_kg*protein/PROTEIN_TO_N_FACTOR/area


def score(rows,key):
    groups=defaultdict(list)
    for r in rows:groups[r['date']].append(r)
    if not groups:return dict(n=0,obs_mean=None,pred_mean=None,RMSE=None,MAE=None,bias=None,NSE=None)
    obs=[sum(float(r['value']) for r in rs)/len(rs) for rs in groups.values()]
    pred=[sum(float(r[key]) for r in rs)/len(rs) for rs in groups.values()]
    mean=sum(obs)/len(obs); residual=[p-o for o,p in zip(obs,pred)]
    sse=sum(e*e for e in residual); sst=sum((o-mean)**2 for o in obs)
    return dict(n=len(obs),obs_mean=mean,pred_mean=sum(pred)/len(pred),RMSE=math.sqrt(sse/len(obs)),
                MAE=sum(abs(e) for e in residual)/len(obs),bias=sum(residual)/len(obs),NSE=1-sse/sst if sst else None)


def fit_terms(pairs,feed):
    groups=defaultdict(list)
    for r in pairs:
        if r['included_in_metrics']=='True' and BACKGROUND_FIT_START<=r['date']<=BACKGROUND_FIT_END:
            groups[r['treatment'],r['date']].append(float(r['value'])-float(r['joint_TN_mg_L']))
    points={t: [(d,sum(v)/len(v)) for (tr,d),v in sorted(groups.items()) if tr==t] for t in ('RM','RCC')}
    if any(len(points[t])<3 for t in points):raise ValueError('每处理至少需要3个背景率定日期')
    background=max(0.,sum(v for d,v in points['RM'])/len(points['RM']))
    xx=sum(feed[d]**2 for d,v in points['RCC'])
    if xx<=0:raise ValueError('RCC背景率定日期没有投饵输入，无法估计投饵项')
    gamma=max(0.,sum(feed[d]*(v-background) for d,v in points['RCC'])/xx)
    return dict(background_mg_L=background,gamma=gamma,
                gamma_unit='(mg/L)/(kg N/ha/day)',training_dates={t:[d for d,v in points[t]] for t in points},
                background_assumption='RCC基础背景与RM相同；RCC剩余平均浓度以投饵经验项表达',
                identifiability='恒定投饵不能区分RCC独有背景和投饵效应；本参数为条件情景估计')


def export_book(path,parameters,metrics,pairs,daily,config,archive):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    book=Workbook(); book.remove(book.active)
    def table(name,rows):
        sheet=book.create_sheet(name)
        if not rows:return
        sheet.append(list(rows[0])); [sheet.append(list(r.values())) for r in rows]
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width=min(65,max(12,max(len(str(c.value or ''))*1.5 for c in col)))
            for c in col[1:]:
                if isinstance(c.value,(int,float)):c.number_format='0.00'
        if name=='说明':
            for row in sheet.iter_rows(min_row=2):
                for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
                sheet.row_dimensions[row[0].row].height=55
    table('参数',[{'参数':'基础背景','数值':parameters['background_mg_L'],'单位':'mg/L','依据':'RM 7月下旬残差均值；RCC共享此背景'},
                  {'参数':'投饵浓度系数gamma','数值':parameters['gamma'],'单位':parameters['gamma_unit'],'依据':'RCC 7月下旬残差，扣除共享背景后估计'},
                  {'参数':'情景日投饵量','数值':FEED_KG_PER_DAY,'单位':'kg/天','依据':'3.0—3.5的中值；采用日记录时以记录为准'},
                  {'参数':'饲料粗蛋白比例','数值':FEED_CRUDE_PROTEIN_FRACTION,'单位':'质量比例','依据':'学位论文印刷页21'},
                  {'参数':'全系统面积','数值':RCC_SYSTEM_AREA_HA,'单位':'ha','依据':'学位论文印刷页21，含虾沟'}])
    table('误差对照',[{'处理':r['treatment'],'范围':r['scope'],'模型':r['model'],'日期数':r['n'],
                      '观测均值_mg_L':r['obs_mean'],'模拟均值_mg_L':r['pred_mean'],'RMSE_mg_L':r['RMSE'],
                      'MAE_mg_L':r['MAE'],'平均偏差_mg_L':r['bias'],'NSE':r['NSE']} for r in metrics])
    table('观测对照',[{'处理':r['treatment'],'日期':r['date'],'观测TN_mg_L':float(r['value']),
                      '施肥模型_mg_L':float(r['joint_TN_mg_L']),'基础背景_mg_L':r['background_mg_L'],
                      '投饵附加项_mg_L':r['feed_contribution_mg_L'],'合计TN_mg_L':r['augmented_TN_mg_L'],
                      '数据用途':r['model_role'],'纳入指标':'是' if r['included_in_metrics']=='True' else '否',
                      '日期依据':r.get('source_date_note',''),'源表行号':int(r['source_row'])} for r in pairs])
    table('逐日组成',daily)
    table('说明',[{'项目':'模型','内容':'C_RM=C_施肥+B；C_RCC=C_施肥+B+gamma×日投饵氮强度。施肥α、k保持已归档值。'},
                  {'项目':'拟合与检验','内容':f'背景和投饵项拟合日期{BACKGROUND_FIT_START}至{BACKGROUND_FIT_END}；{HOLDOUT_START}起观测留出。全期指标含前后两组率定数据。'},
                  {'项目':'投饵情景','内容':f'{FEED_START}至{FEED_END}连续投喂{FEED_KG_PER_DAY} kg/天为情景假设。论文未提供逐日开停喂记录。'},
                  {'项目':'识别边界','内容':parameters['identifiability']+'；不能据此作因果归因或独立氮收支解释。'},
                  {'项目':'文献边界','内容':'学位论文式4-15的投饵线性项原用于RCI，本版扩展到RCC。蛋白/6.25估算氮含量，假设与投喂量质量基准一致。'},
                  {'项目':'区域输入','内容':'施肥模型仍使用Nfer-2018作为2022稻季代理及SOM百分数假设；背景/投饵参数为点位经验值。'},
                  {'项目':'完整记录','内容':str(archive)}])
    book.save(path)


def make_plot(path,daily,pairs,metrics):
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.font_manager import FontProperties
    from matplotlib.lines import Line2D
    from matplotlib import rc_context
    import matplotlib.dates as mdates
    font=FontProperties(fname='C:/Windows/Fonts/msyh.ttc').get_name()
    with rc_context({'font.family':font,'axes.unicode_minus':False,'font.size':10}):
        fig=Figure(figsize=(14,8)); FigureCanvasAgg(fig)
        axes=fig.subplots(2,2,gridspec_kw={'width_ratios':[1.6,1]})
        fig.subplots_adjust(left=.065,right=.985,bottom=.12,top=.84,hspace=.34,wspace=.22)
        for row,t in enumerate(('RM','RCC')):
            obs=[r for r in pairs if r['treatment']==t]
            for col,ax in enumerate(axes[row]):
                left,right=('2022-06-01','2022-10-31') if col==0 else ('2022-08-01','2022-09-05')
                days=[r for r in daily if r['处理']==t and left<=r['日期']<=right]
                x=[date.fromisoformat(r['日期']) for r in days]
                ax.plot(x,[r['施肥响应_mg_L'] for r in days],color='#999999',ls='--',lw=1.5)
                ax.plot(x,[r['合计TN_mg_L'] for r in days],color='#177e96',lw=2)
                for accepted,marker,color in [(True,'o','#222222'),(False,'x','#bd803b')]:
                    points=[r for r in obs if (r['included_in_metrics']=='True')==accepted and left<=r['date']<=right]
                    ax.scatter([date.fromisoformat(r['date']) for r in points],[float(r['value']) for r in points],marker=marker,color=color,s=27,zorder=4)
                ax.set_xlim(date.fromisoformat(left),date.fromisoformat(right)); ax.set_ylim(0,30 if col==0 else 2.6)
                ax.spines[['top','right']].set_visible(False); ax.grid(alpha=.15); ax.set_ylabel('TN（mg/L）')
                ax.xaxis.set_major_locator(mdates.MonthLocator() if col==0 else mdates.DayLocator(interval=7)); ax.xaxis.set_major_formatter(mdates.DateFormatter('%m-%d'))
                ax.set_title(t+('｜6—10月模拟' if col==0 else '｜8月以后留出检验'),loc='left',fontsize=12)
                if col==0:
                    ax.axvspan(date.fromisoformat(BACKGROUND_FIT_START),date.fromisoformat(BACKGROUND_FIT_END),color='#e8f1ed',zorder=0)
                    last=max(date.fromisoformat(r['date']) for r in obs)
                    ax.axvspan(last,date.fromisoformat(right),color='#eeeeee',zorder=0)
                    ax.text(.8,.7,'观测结束后\n情景预测',transform=ax.transAxes,ha='center',color='#666666')
                else:
                    old=next(r for r in metrics if r['treatment']==t and r['scope']=='8月以后留出' and r['model']=='施肥模型')
                    new=next(r for r in metrics if r['treatment']==t and r['scope']=='8月以后留出' and r['model']=='背景与投饵情景')
                    ax.text(.98,.97,f"n={new['n']}\nRMSE：{old['RMSE']:.2f} → {new['RMSE']:.2f} mg/L\n观测/模拟均值：{new['obs_mean']:.2f} / {new['pred_mean']:.2f}",transform=ax.transAxes,ha='right',va='top',fontsize=9,bbox=dict(facecolor='white',edgecolor='none',alpha=.85))
        fig.suptitle('2022年田面水TN：背景浓度与投饵经验项情景',fontsize=17,y=.98)
        handles=[Line2D([],[],color='#999999',ls='--',label='施肥响应'),Line2D([],[],color='#177e96',lw=2,label='加入背景与投饵项'),
                 Line2D([],[],color='#222222',marker='o',ls='',label='明确日期观测'),Line2D([],[],color='#bd803b',marker='x',ls='',label='日期/顺序待核观测')]
        fig.legend(handles=handles,loc='upper center',bbox_to_anchor=(.52,.936),ncol=4,frameon=False)
        fig.text(.5,.045,'浅绿：背景与投饵项率定期；灰色：超出观测覆盖。投饵按3.25 kg/天情景，共享基础背景假设；左右纵轴不同。',ha='center',fontsize=10)
        fig.savefig(path,dpi=200)


def main():
    source_pairs=read(INPUT_ARCHIVE/'observation_pairs.csv'); source_daily=read(INPUT_ARCHIVE/'daily_comparison.csv')
    dates=sorted({r['日期'] for r in source_daily})
    feed_records={r['date']:float(r['feed_kg']) for r in read(Path(FEED_DAILY_CSV))} if FEED_DAILY_CSV else None
    if feed_records is not None and set(dates)-feed_records.keys():raise ValueError('投饵CSV缺少模拟日期；停喂日期请明确填写0')
    feed_kg={d:(feed_records[d] if feed_records is not None else (FEED_KG_PER_DAY if FEED_START<=d<=FEED_END else 0.)) for d in dates}
    feed_n={d:nitrogen_from_feed(v) for d,v in feed_kg.items()}
    parameters=fit_terms(source_pairs,feed_n)
    daily=[]; lookup={}
    for r in source_daily:
        t,d=r['处理'],r['日期']; fert=float(r['联合校准TN_mg_L'])
        fn=feed_n[d] if t=='RCC' else 0.; contribution=parameters['gamma']*fn
        item={'处理':t,'日期':d,'施肥响应_mg_L':fert,'基础背景_mg_L':parameters['background_mg_L'],
              '日投饵量_kg':feed_kg[d] if t=='RCC' else 0.,'日投饵氮_kgN_ha':fn,
              '投饵附加项_mg_L':contribution,'合计TN_mg_L':fert+parameters['background_mg_L']+contribution}
        daily.append(item); lookup[t,d]=item
    pairs=[]
    for source in source_pairs:
        r=dict(source); item=lookup[r['treatment'],r['date']]
        role='其他日期'
        if r['role_in_trial']=='首次事件率定':role='施肥峰值率定'
        elif r['date'] in parameters['training_dates'][r['treatment']]:role='背景与投饵项率定'
        elif r['date']>=HOLDOUT_START:role='后期留出'
        if r['included_in_metrics']!='True':role+='；日期或采样顺序待核，仅图示'
        r.update(background_mg_L=item['基础背景_mg_L'],feed_contribution_mg_L=item['投饵附加项_mg_L'],augmented_TN_mg_L=item['合计TN_mg_L'],model_role=role)
        pairs.append(r)
    metrics=[]
    for t in ('RM','RCC'):
        for scope in ('背景与投饵项率定','8月以后留出','全部日期（含两类率定）'):
            selected=[r for r in pairs if r['treatment']==t and r['included_in_metrics']=='True']
            if scope=='背景与投饵项率定':selected=[r for r in selected if r['date'] in parameters['training_dates'][t]]
            if scope=='8月以后留出':selected=[r for r in selected if r['date']>=HOLDOUT_START]
            for name,key in [('施肥模型','joint_TN_mg_L'),('背景与投饵情景','augmented_TN_mg_L')]:
                metrics.append(dict(treatment=t,scope=scope,model=name,**score(selected,key)))
    run_id=datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    archive=HERE/'feed_runs'/run_id; archive.mkdir(parents=True)
    config={key:str(value) if isinstance(value,Path) else value for key,value in globals().items() if key.isupper() and key not in ('HERE','ROOT')}
    config['method']='fixed fertilizer response + shared RM background + gamma * daily feed N per system hectare'
    config['sources']=['张怡杰学位论文印刷页21、54，式4-15（RCI扩展为RCC情景）','https://www.fao.org/4/s4314e/s4314e0v.htm']
    for name,rows in [('daily_components',daily),('observation_pairs',pairs),('metrics',metrics)]:
        with (archive/(name+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:
            writer=csv.DictWriter(f,fieldnames=list(rows[0])); writer.writeheader(); writer.writerows(rows)
    for name,obj in [('config_snapshot',config),('parameters',parameters)]:
        (archive/(name+'.json')).write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')
    shutil.copy2(__file__,archive/Path(__file__).name)
    for name in ('observation_source.xlsx','joint_calibration.json'):
        shutil.copy2(INPUT_ARCHIVE/name,archive/name)
    if FEED_DAILY_CSV:shutil.copy2(FEED_DAILY_CSV,archive/'feed_source.csv')
    public=OUTPUT_DIRECTORY/('2022背景投饵情景_'+run_id); public.mkdir(parents=True)
    export_book(public/'2022背景投饵结果.xlsx',parameters,metrics,pairs,daily,config,archive)
    make_plot(public/'TN背景投饵对照.png',daily,pairs,metrics)
    manifest=dict(public_output=str(public),source_archive=str(INPUT_ARCHIVE),
                  source_sha256={name:sha(INPUT_ARCHIVE/name) for name in ('observation_pairs.csv','daily_comparison.csv','joint_calibration.json')},
                  output_sha256={p.name:sha(p) for p in public.iterdir()},code_sha256=sha(Path(__file__)))
    (archive/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(parameters,ensure_ascii=False,indent=2)); print(json.dumps([r for r in metrics if r['scope']=='8月以后留出'],ensure_ascii=False,indent=2))
    print('结果：'+str(public)); print('记录：'+str(archive))
    return archive


if __name__ in ('__main__','__console__'):main()
