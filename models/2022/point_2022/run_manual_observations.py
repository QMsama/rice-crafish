"""以人工点选CSV重跑既定点位浓度模型；保留源记录及筛选依据。"""
from pathlib import Path
from datetime import date, datetime, timedelta
import csv,json,math,shutil,sys,hashlib,copy
sys.dont_write_bytecode=True
HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
sys.path.insert(0,str(HERE))
import run_joint_point_2022 as joint
import run_feed_background_2022 as feed
SOURCE=ROOT/'模型/2026.9.3/2022点位验证/2022田面水TN_人工选点_20261018_202627_196311.csv'
FROZEN=HERE/'joint_runs/20261008_221451_815334'
# 只有得到用户确认后，才在此登记日期更正；原CSV保持原样。
DATE_CORRECTIONS={'RM-02':'2022-06-13','RM-37':'2022-09-02'}
DATE_REVIEW_IDS={'RM-02','RM-37'}
def write(path,rows):
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
def predict(t,d,events,parameters=None):
    own=sorted([e for e in events if e['treatment']==t and e['response_date']<=d],key=lambda e:e['response_date'])
    if not own:return 0.
    e=own[-1];age=(date.fromisoformat(d)-date.fromisoformat(e['response_date'])).days
    alpha=parameters[t]['alpha'] if parameters else 1.
    k=parameters[t]['k'] if parameters else float(e['k_day'])
    return alpha*float(e['C0_mg_L'])*math.exp(-k*age)
def main():
    raw=feed.read(SOURCE);events=feed.read(FROZEN/'event_parameters.csv')
    assert len({r['点编号'] for r in raw})==len(raw)
    assert len(raw)==76
    pairs=[]
    for rownum,r in enumerate(raw,2):
        t=r['处理'];assert t in ('RM','RCC')
        d=DATE_CORRECTIONS.get(r['点编号'],r['日期']);date.fromisoformat(d)
        assert '2022-06-01'<=d<='2022-10-31'
        value=float(r['TN_mg_L']);assert math.isfinite(value) and value>=0
        own=[e for e in events if e['treatment']==t];reasons=[]
        if r['日期状态']!='用户确认':reasons.append('日期待确认')
        if r['点编号'] in DATE_REVIEW_IDS and r['点编号'] not in DATE_CORRECTIONS:reasons.append('录入日期与原图点序不一致待确认')
        if d in [e['date'] for e in own]:reasons.append('施肥当天采样先后待确认')
        if d<min(e['response_date'] for e in own):reasons.append('施肥响应前初始状态')
        pairs.append(dict(record_id=r['点编号'],treatment=t,date=d,original_date=r['日期'],value=value,
            included_in_metrics=str(not reasons),evaluation_status='；'.join(reasons) or '纳入',
            source_row=rownum,source_file=str(SOURCE),source_date_note='用户确认更正' if r['点编号'] in DATE_CORRECTIONS else r['日期状态'],
            baseline_TN_mg_L=predict(t,d,events),joint_TN_mg_L=0.))
    parameters={t:joint.fit_treatment(events,pairs,t) for t in ('RM','RCC')}
    for r in pairs:
        r['joint_TN_mg_L']=predict(r['treatment'],r['date'],events,parameters)
        r['role_in_trial']='首次事件率定' if r['included_in_metrics']=='True' and r['date'] in parameters[r['treatment']]['training_dates'] else '未参与率定'
    days=[(date(2022,6,1)+timedelta(days=i)).isoformat() for i in range(153)]
    fn={d:feed.nitrogen_from_feed(feed.FEED_KG_PER_DAY) for d in days}
    bg=feed.fit_terms(pairs,fn)
    daily=[];lookup={}
    for t in ('RM','RCC'):
        for d in days:
            fert=predict(t,d,events,parameters);n=fn[d] if t=='RCC' else 0.;contribution=bg['gamma']*n
            item={'处理':t,'日期':d,'施肥响应_mg_L':fert,'基础背景_mg_L':bg['background_mg_L'],
                '日投饵量_kg':feed.FEED_KG_PER_DAY if t=='RCC' else 0.,'日投饵氮_kgN_ha':n,
                '投饵附加项_mg_L':contribution,'合计TN_mg_L':fert+bg['background_mg_L']+contribution}
            daily.append(item);lookup[t,d]=item
    for r in pairs:
        item=lookup[r['treatment'],r['date']]
        role='其他评价日期'
        if r['role_in_trial']=='首次事件率定':role='施肥峰值率定'
        elif r['date'] in bg['training_dates'][r['treatment']]:role='背景与投饵项率定'
        elif r['date']>='2022-08-01':role='后期留出'
        r.update(background_mg_L=bg['background_mg_L'],feed_contribution_mg_L=item['投饵附加项_mg_L'],augmented_TN_mg_L=item['合计TN_mg_L'],model_role=role if r['included_in_metrics']=='True' else r['evaluation_status'])
    metrics=[]
    for t in ('RM','RCC'):
        second=sorted(e['response_date'] for e in events if e['treatment']==t)[1]
        scopes={'首次事件率定':lambda r:r['date'] in parameters[t]['training_dates'],
            '第二次事件前7天留出':lambda r:0<=(date.fromisoformat(r['date'])-date.fromisoformat(second)).days<7,
            '背景与投饵项率定':lambda r:r['date'] in bg['training_dates'][t],
            '8月以后留出':lambda r:r['date']>='2022-08-01','全部日期（含两类率定）':lambda r:True}
        for scope,filter_ in scopes.items():
            selected=[r for r in pairs if r['treatment']==t and r['included_in_metrics']=='True' and filter_(r)]
            for name,key in [('施肥模型','joint_TN_mg_L'),('背景与投饵情景','augmented_TN_mg_L')]:
                metrics.append(dict(treatment=t,scope=scope,model=name,**feed.score(selected,key)))
    stamp=datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    archive=HERE/'manual_runs'/stamp;archive.mkdir(parents=True)
    public=feed.OUTPUT_DIRECTORY/('2022人工观测重算_'+stamp);public.mkdir()
    corrected=copy.deepcopy(raw); counts={}
    for row in corrected:
        row['日期']=DATE_CORRECTIONS.get(row['点编号'],row['日期'])
        key=(row['处理'],row['日期']);counts[key]=counts.get(key,0)+1
        row['同日点序号']=str(counts[key])
    write(public/'2022田面水TN_人工选点_日期确认.csv',corrected)
    for name,data in [('observations',pairs),('daily_components',daily),('metrics',metrics)]:write(archive/(name+'.csv'),data)
    settings={'source':str(SOURCE),'frozen_inputs':str(FROZEN),'date_corrections':DATE_CORRECTIONS,'pending_ids':sorted(DATE_REVIEW_IDS-set(DATE_CORRECTIONS)),
        'peak_parameters':parameters,'background_feed_parameters':bg,'feed_kg_day':feed.FEED_KG_PER_DAY,'protein_fraction':feed.FEED_CRUDE_PROTEIN_FRACTION,'system_area_ha':feed.RCC_SYSTEM_AREA_HA,
        'source_sha256':feed.sha(SOURCE),'event_sha256':feed.sha(FROZEN/'event_parameters.csv'),'public_output':str(public)}
    (archive/'config.json').write_text(json.dumps(settings,ensure_ascii=False,indent=2),encoding='utf-8')
    shutil.copy2(SOURCE,archive/'observation_source.csv');shutil.copy2(FROZEN/'event_parameters.csv',archive/'event_parameters.csv')
    for p in [Path(__file__),Path(joint.__file__),Path(feed.__file__)]:shutil.copy2(p,archive/p.name)
    feed.export_book(public/'2022人工观测重算结果.xlsx',bg,metrics,pairs,daily,settings,archive)
    from openpyxl import load_workbook
    book=load_workbook(public/'2022人工观测重算结果.xlsx')
    book['说明']['B2']='C_RM=C_施肥+B；C_RCC=C_施肥+B+gamma×日投饵氮强度。本轮使用人工观测重新率定α、k、B、gamma。'
    book['说明'].append(['观测来源与筛选','输入为人工点选CSV；日期与图中点序冲突、施肥当天先后待核及初始状态记录单独标记，不参与指标。源表行号指CSV行号。'])
    s=book.create_sheet('峰值参数');s.append(['处理','alpha','k_每天','k下限','k上限','训练日期数','训练日期'])
    for t,p in parameters.items():s.append([t,p['alpha'],p['k'],p['k_lower'],p['k_upper'],p['n_dates'],'、'.join(p['training_dates'])])
    for row in s.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value,(int,float)):c.number_format='0.00'
    s.column_dimensions['G'].width=80
    book.save(public/'2022人工观测重算结果.xlsx')
    # 本次橙叉表示所有未纳入指标的记录，含初始状态，明确图例含义。
    import inspect
    plot_code=inspect.getsource(feed.make_plot).replace("label='明确日期观测'","label='用户确认且纳入评价'").replace("label='日期/顺序待核观测'","label='未纳入评价的观测'")
    context=dict(feed.__dict__);exec(plot_code,context)
    context['make_plot'](public/'TN人工观测重算对照.png',daily,pairs,metrics)
    write(public/'2022人工观测与模拟对照.csv',[{'点编号':r['record_id'],'处理':r['treatment'],'原录入日期':r['original_date'],'计算日期':r['date'],'观测TN_mg_L':f"{r['value']:.2f}",'模拟TN_mg_L':f"{r['augmented_TN_mg_L']:.2f}",'评价用途':r['model_role'],'筛选说明':r['evaluation_status']} for r in pairs])
    # 独立优化复核及留出隔离检查。
    from scipy.optimize import least_squares
    for t,p in parameters.items():
        first=min((e for e in events if e['treatment']==t),key=lambda e:e['date']);start=date.fromisoformat(first['response_date'])
        selected=[r for r in pairs if r['treatment']==t and r['included_in_metrics']=='True' and r['date'] in p['training_dates']]
        fit=least_squares(lambda x:[x[0]*float(first['C0_mg_L'])*math.exp(-x[1]*(date.fromisoformat(r['date'])-start).days)-r['value'] for r in selected],x0=[p['alpha'],(p['k_lower']+p['k_upper'])/2],bounds=([0,p['k_lower']],[10,p['k_upper']]),gtol=1e-12,ftol=1e-12,xtol=1e-12)
        assert abs(fit.x[0]-p['alpha'])<1e-5 and abs(fit.x[1]-p['k'])<1e-5
    changed=copy.deepcopy(pairs)
    for r in changed:
        if r['date']>='2022-08-01':r['value']=10000
    assert feed.fit_terms(changed,fn)==bg
    assert {t:joint.fit_treatment(events,changed,t) for t in parameters}==parameters
    assert len(daily)==306 and len(pairs)==76
    (archive/'verification.json').write_text(json.dumps({'independent_optimizer':'PASS','heldout_isolation':'PASS','daily_rows':306,'observation_rows':76,'outputs':{p.name:feed.sha(p) for p in public.iterdir()}},ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(settings,ensure_ascii=False,indent=2));print(json.dumps([r for r in metrics if r['scope']=='8月以后留出'],ensure_ascii=False,indent=2))
    return archive
if __name__=='__main__':main()
