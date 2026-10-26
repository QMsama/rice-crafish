"""2022田面水深基准情景：现有T分量加明确的补充蒸发假设。"""
from pathlib import Path
from datetime import date,datetime
import sys,csv,json,shutil,hashlib,math,os
sys.dont_write_bytecode=True
HERE=Path(__file__).resolve().parent; ROOT=HERE.parents[1]
os.environ['MPLCONFIGDIR']=str(HERE/'cache/matplotlib')
sys.path.insert(0,str(ROOT/'模型/2026.6.27/python计算'))
from model_v6.regional_hydrology import step_rm_grid,LossParameters,WaterLevelRule
from model_v6.management import rm_water_rule,rcc_water_rule,is_managed_drydown_transition,INACTIVE
from openpyxl import load_workbook

# 可修改参数：均为情景输入，尚未按2022田块实际灌排记录核定。
SUPPLEMENTAL_EVAPORATION_MM_DAY=2.0
SEEPAGE_MM_DAY=2.0
DRY_STAGE_ABSORPTION_MM_DAY=60.0
INITIAL_DEPTH_MM={'RM':0.0,'RCC':0.0}
WEATHER=HERE/'run_records/20261008_221451_861261/point_weather.csv'
OBSERVATIONS=ROOT/'模型/2026.9.3/原始数据/2022年大田实验数据/2022年大田实验数据提取.xlsx'
OUTPUT=ROOT/'模型/2026.9.3/2022点位验证'

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def read(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def write(p,rows,rounded=False):
    with p.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader()
        for r in rows:w.writerow({k:f'{v:.2f}' if rounded and isinstance(v,float) else v for k,v in r.items()})
def main():
    weather=read(WEATHER)
    assert len(weather)==153 and len({r['date'] for r in weather})==153
    for r in weather:
        for key in ('rainfall_mm_day','transpiration_mm_day'):assert math.isfinite(float(r[key])) and float(r[key])>=0
    losses=LossParameters(SEEPAGE_MM_DAY,SUPPLEMENTAL_EVAPORATION_MM_DAY,DRY_STAGE_ABSORPTION_MM_DAY)
    daily=[];lookup={}
    for t in ('RM','RCC'):
        depth=INITIAL_DEPTH_MM[t];previous=INACTIVE
        for r in weather:
            d=date.fromisoformat(r['date']);rule=rm_water_rule(d) if t=='RM' else rcc_water_rule(d)
            initial=depth;p=float(r['rainfall_mm_day']);tr=float(r['transpiration_mm_day'])
            result=step_rm_grid(initial_depth_mm=depth,rainfall_mm=p,transpiration_mm=tr,rule=rule,losses=losses,managed_drydown=is_managed_drydown_transition(previous,rule))
            depth=float(result.final_depth_mm);irrig=float(result.irrigation_mm);drain=float(result.boundary_drain_mm)
            loss=float(result.actual_surface_loss_mm);absorb=float(result.dry_stage_absorption_mm)
            # 独立重算水量闭合；RCC出流仅表示田面离水，未建虾沟出口。
            residual=initial+p+irrig-loss-absorb-drain-depth
            assert abs(residual)<1e-9 and depth>=0
            row={'处理':t,'日期':r['date'],'期初水深_mm':initial,'模拟田面水深_mm':depth,'降雨_mm':p,'蒸腾T_mm':tr,
                '补充蒸发假设_mm':SUPPLEMENTAL_EVAPORATION_MM_DAY,'潜在下渗_mm':SEEPAGE_MM_DAY,
                '实际合计耗水_mm':loss,'情景灌溉_mm':irrig,'田面出流_mm':drain,'干田吸收_mm':absorb,
                '灌水触发水深_mm':float(rule.minimum_mm),'灌水目标水深_mm':float(rule.target_mm),'最大蓄水深_mm':float(rule.maximum_mm),
                '水量闭合差_mm':residual,'出流边界':'田面出口' if t=='RM' else '田面至虾沟，系统外排未计算'}
            daily.append(row);lookup[t,r['date']]=depth;previous=rule
    book=load_workbook(OBSERVATIONS,read_only=True,data_only=True);sheet=book['田面水位动态变化'];pairs=[]
    for n,values in enumerate(sheet.iter_rows(min_row=2,values_only=True),2):
        d,t,name,seq,cm,mm,status=values
        if t not in INITIAL_DEPTH_MM:continue
        d=d.date().isoformat();observed=float(mm)
        assert abs(float(cm)*10-observed)<.02
        pred=lookup[t,d]
        pairs.append({'处理':t,'参考日期':d,'同日序号':seq or '', '观测水深_mm':observed,'情景水深_mm':pred,
            '差值_mm':pred-observed,'日期依据':status,'源工作表行号':n,'评价用途':'日期及采样时刻待核，仅初步对照'})
    book.close()
    stamp=datetime.now().strftime('%Y%m%d_%H%M%S_%f');archive=HERE/'water_runs'/stamp;archive.mkdir(parents=True)
    public=OUTPUT/('2022田面水深情景_'+stamp);public.mkdir()
    write(archive/'daily_full_precision.csv',daily);write(archive/'observation_pairs.csv',pairs)
    write(public/'2022逐日田面水量平衡.csv',daily,True);write(public/'2022水位观测对照.csv',pairs,True)
    manifest={'public_output':str(public),'model':'daily surface storage baseline; no observation calibration',
        'formula':'h[t]=h[t-1]+P+I-actual_loss-dry_absorption-field_outflow',
        'assumptions':{'supplemental_E_mm_day':SUPPLEMENTAL_EVAPORATION_MM_DAY,'seepage_mm_day':SEEPAGE_MM_DAY,
            'dry_absorption_mm_day':DRY_STAGE_ABSORPTION_MM_DAY,'initial_depth_mm':INITIAL_DEPTH_MM,
            'management':'model_v6.management regional scenario calendar, not confirmed 2022 field operations',
            'sampling':'end of day simulation; within-day sampling order unconfirmed'},
        'scope':'RM and RCC field depth only; no RCC ditch boundary flow or nitrogen load',
        'sources':{'weather_sha256':sha(WEATHER),'workbook_sha256':sha(OBSERVATIONS)},
        'ET_product':'https://www.tpdc.ac.cn/en/data/560b250c-43d6-4bd9-98db-da97b3556f30',
        'literature':'张怡杰学位论文印刷页61-62式5-1至5-6、页65田面与养殖沟区分',
        'daily_rows':len(daily),'observations':len(pairs),'max_abs_balance_error_mm':max(abs(r['水量闭合差_mm']) for r in daily),
        'formal_validation_dates':0,'diagnostic_metrics':[]}
    for t in INITIAL_DEPTH_MM:
        group=[r for r in pairs if r['处理']==t];bydate={}
        for r in group:bydate.setdefault(r['参考日期'],[]).append(r)
        obs=[sum(r['观测水深_mm'] for r in group)/len(group) for group in bydate.values()]
        pred=[group[0]['情景水深_mm'] for group in bydate.values()]
        rmse=math.sqrt(sum((a-b)**2 for a,b in zip(obs,pred))/len(obs))
        manifest['diagnostic_metrics'].append({'treatment':t,'estimated_date_count':len(obs),'RMSE_mm':rmse,'purpose':'仅按估计日期日均值诊断，不作为正式验证'})
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib import rc_context
    from matplotlib.font_manager import FontProperties
    import matplotlib.dates as mdates
    with rc_context({'font.family':FontProperties(fname='C:/Windows/Fonts/msyh.ttc').get_name(),'axes.unicode_minus':False}):
        fig=Figure(figsize=(13,8));FigureCanvasAgg(fig);axes=fig.subplots(2,1)
        fig.subplots_adjust(left=.08,right=.97,bottom=.15,top=.88,hspace=.36)
        for ax,t in zip(axes,INITIAL_DEPTH_MM):
            rs=[r for r in daily if r['处理']==t];ps=[r for r in pairs if r['处理']==t]
            ax.plot([date.fromisoformat(r['日期']) for r in rs],[r['模拟田面水深_mm'] for r in rs],color='#177e96',lw=1.7,label='规则驱动的日末田面水深')
            ax.scatter([date.fromisoformat(r['参考日期']) for r in ps],[r['观测水深_mm'] for r in ps],marker='x',s=30,color='#b17b3f',label='图件观测（日期待核）')
            ax.set_title(t,loc='left');ax.set_ylabel('田面水深（mm）');ax.set_xlim(date(2022,6,1),date(2022,10,31));ax.set_ylim(bottom=0)
            ax.xaxis.set_major_locator(mdates.MonthLocator());ax.xaxis.set_major_formatter(mdates.DateFormatter('%m-%d'));ax.grid(alpha=.15)
            last=max(date.fromisoformat(r['参考日期']) for r in ps);ax.axvspan(last,date(2022,10,31),color='#eeeeee',zorder=0)
            ax.text(.78,.85,'观测覆盖期之后',transform=ax.transAxes,color='#666666');ax.legend(loc='upper left',bbox_to_anchor=(.30,1.15),ncol=2,frameon=False,fontsize=9)
        fig.suptitle('2022年田面水深基准情景',fontsize=17)
        fig.text(.5,.07,'输入：降雨＋蒸腾T；补充蒸发2 mm/天、下渗2 mm/天为假设。灌排日程沿用区域情景规则。',ha='center',fontsize=10)
        fig.text(.5,.035,'观测日期为插值/外推；本图用于诊断。RCC田面出流不等于种养系统对外排水。',ha='center',fontsize=10)
        fig.savefig(public/'2022田面水深对照.png',dpi=180)
    shutil.copy2(WEATHER,archive/'point_weather.csv');shutil.copy2(OBSERVATIONS,archive/'observation_source.xlsx');shutil.copy2(__file__,archive/Path(__file__).name)
    for name in ['management.py','regional_hydrology.py']:shutil.copy2(ROOT/'模型/2026.6.27/python计算/model_v6'/name,archive/name)
    manifest['output_sha256']={p.name:sha(p) for p in public.iterdir()}
    (archive/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(manifest,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
