"""2022年潜江点位TN浓度计算。

日常修改：只需编辑下面的“一、用户参数区”，保存后在QGIS重新运行。
计算逻辑：施肥事件 → 初始浓度C0 → 逐日衰减 → 与观测配对 → 导出结果。
每次运行保存实际参数快照；所有日期统一写为 YYYY-MM-DD。
"""
from pathlib import Path
from datetime import date, datetime, timedelta
import csv
import hashlib
import json
import math
import os
import sys
import copy
import runpy

sys.dont_write_bytecode = True
HERE = Path(__file__).resolve().parent
os.environ['MPLCONFIGDIR'] = str(HERE/'cache'/'matplotlib')

# ============================================================================
# 一、用户参数区：路径、日期、点位和施肥事件均在这里修改
# ============================================================================

# 1. 项目根目录。Windows路径建议使用 r"..." 或正斜杠，避免反斜杠转义。
PROJECT_ROOT = r"E:/#岑俊毅/2.论文/2.稻虾综合种养"

# 2. 输入路径：支持绝对路径，也支持相对于 PROJECT_ROOT 的路径。
OBSERVATION_FILE = "模型/2026.9.3/原始数据/2022年大田实验数据/2022年大田实验数据提取.xlsx"
OBSERVATION_SHEET = "田面水TN浓度变化"  # 每次运行直接读取此表的最新保存值
RAIN_DIRECTORY = "模型/2026.9.3/模型背景数据/rain daily tiff/2022"
TRANSPIRATION_DIRECTORY = "模型/2026.9.3/模型背景数据/蒸散发/T_2022_daily_tiff"
SOIL_INTERPOLATION_FILE = ".codex-work/soil-source-check-20261004/interpolation.json"

# 施氮栅格：在点位原始网格中取值，先按栅格scale/offset解码。
NFER_RASTER = "模型/2026.9.3/模型背景数据/nitrogen_fertilizer_use/Nfer-2018"
NFER_UNIT_FACTOR = 10.0  # g N/m²土地/年 → kg N/ha土地/年
# 下列系数是年度土地口径到稻季田块口径的情景映射，不能由月份比例推定。
# 1.0表示直接采用年度土地强度作为稻季代理；尚未完成作物/面积校正。
NFER_TO_RICE_SEASON_FACTOR = 1.0
NFER_MAPPING_STATUS = "annual_land_intensity_as_rice_season_proxy; crop_area_mapping_pending"

# 若项目搬家，修改这两个土壤栅格路径即可；仍会校验其内容与插值来源一致。
SOIL_SOURCE_FILES = {
    "pH": "模型/2026.9.3/模型背景数据/soil/Environmental Date_pH.tif",
    "SOM": "模型/2026.9.3/模型背景数据/soil/Environmental Data_SOM.tif",
}

# 3. 最终结果输出父目录；脚本和缓存保存在项目 .codex-work 中。
# 每次运行在这里新建子目录，提供Excel和对照图；复现记录自动存入.codex-work。
OUTPUT_DIRECTORY = r"E:/#岑俊毅/2.论文/2.稻虾综合种养/模型/2026.9.3/2022点位验证"
LOAD_TABLES_IN_QGIS = False  # True：同时加载归档CSV表；False：查看Excel和对照图

# 4. 本脚本用于2022年验证；可在2022年内调整起止日期，包含起止当天。
START_DATE = "2022-06-01"
END_DATE = "2022-10-31"
LONGITUDE = 112.6166  # 东经，十进制度
LATITUDE = 30.3833   # 北纬，十进制度
# 修改点位后，需要同时更换与新坐标对应的土壤插值文件。

# 5. 栅格文件命名规则：year=年份，doy=年内日序（001—365）。
RAIN_FILENAME_PATTERN = "PRE_{year}_Day_{doy:03d}.tif"
TRANSPIRATION_FILENAME_PATTERN = "T_{year}{doy:03d}.tif"

# 6. 土壤与浓度设置。
SOM_SOURCE_UNIT = "percent"  # percent=百分数；g/kg=克/千克（自动除以10）
SOM_UNIT_CONFIRMED = False  # 原始单位说明核实后再改为True；本项用于记录证据状态
INITIAL_TN_MG_L = 0.0      # 首次施肥前的启动浓度，mg/L；该时段暂不计入评价
EVENT_RULE = "latest_event"  # latest_event=每次施肥重设C0；sum_events=各次贡献叠加
RUN_LABEL = "climate2022_nfer2018_raster_proxy"  # 记录气象年与施氮数据年的组合
CALIBRATE_K = True  # 用首次施肥后7天内、日期明确的观测率定k；C0仍由栅格计算
K_BOUNDS = {"basal": [0.213, 0.436], "tillering": [0.231, 0.512], "panicle": [0.263, 0.670]}

# 7. 逐次施肥事件：一对大括号代表一次施肥，可复制整段增加事件。
# treatment：RM=水稻单作，RCC=稻虾共作。
# date：施肥日期；N_fraction：本次占该处理稻季施氮代理量的比例。
# 每个处理全部事件的N_fraction之和必须为1；N_kg_ha由栅格自动计算。
# stage：basal=基肥，tillering=蘖肥，panicle=穗肥；决定使用哪组文献系数。
# response_lag_days：施肥日到C0对应日的间隔；1表示施肥后第1天为C0。
# stage_status记录阶段依据；默认分配比例依据论文6:4日程。
FERTILIZER_EVENTS = [
    {
        "treatment": "RM",
        "date": "2022-06-18",
        "stage": "tillering",
        "N_fraction": 0.60,
        "response_lag_days": 1,
        "stage_status": "manuscript_蘖肥",
    },
    {
        "treatment": "RM",
        "date": "2022-07-06",
        "stage": "tillering",
        "N_fraction": 0.40,
        "response_lag_days": 1,
        "stage_status": "manuscript_追肥; tillering_mapping_pending",
    },
    {
        "treatment": "RCC",
        "date": "2022-06-14",
        "stage": "basal",
        "N_fraction": 0.60,
        "response_lag_days": 1,
        "stage_status": "manuscript_底肥",
    },
    {
        "treatment": "RCC",
        "date": "2022-06-27",
        "stage": "tillering",
        "N_fraction": 0.40,
        "response_lag_days": 1,
        "stage_status": "manuscript_蘖肥",
    },
]

# 8. 文献系数：C0 = a0 + aN×N + apH×pH + aSOM×SOM；k_day单位为天^-1。
# 依据Ruan等2022表2、表3长江流域参数；通常保留，研究调整时记录依据。
COEFFICIENTS = {
    "basal":     {"a0": -21.469,  "aN": 0.841, "apH": 3.074,  "aSOM": 0.0,     "k_day": 0.370},
    "tillering": {"a0": -101.200, "aN": 0.540, "apH": 22.184, "aSOM": -11.093, "k_day": 0.339},
    "panicle":   {"a0": -164.918, "aN": 0.435, "apH": 26.256, "aSOM": 0.0,     "k_day": 0.501},
}

# 9. 来源文件：用于生成复现校验值；更换资料时修改相应路径。
SOURCE_FILES = [
    "Rice-crayfish systems 20260920 V1.docx",
    ".codex-work/model-literature/1-s2.0-S0301479722005059-main.txt",
    ".codex-work/model-literature/张怡杰-学位论文.txt",
    "模型/2026.9.3/原始数据/2022年大田实验数据/2022年大田实验数据提取.xlsx",
]

# ============================================================================
# 二、参数整理与检查：日常修改集中在上方，以下为计算函数
# ============================================================================


def build_config():
    """将上方用户参数整理为独立快照，供计算和结果归档使用。"""
    return copy.deepcopy({
        "project_root": PROJECT_ROOT,
        "output_directory": OUTPUT_DIRECTORY,
        "start_date": START_DATE, "end_date": END_DATE,
        "longitude": LONGITUDE, "latitude": LATITUDE,
        "observations": OBSERVATION_FILE,
        "calibrate_k": CALIBRATE_K, "k_bounds": K_BOUNDS,
        "observation_sheet": OBSERVATION_SHEET,
        "rain_directory": RAIN_DIRECTORY,
        "transpiration_directory": TRANSPIRATION_DIRECTORY,
        "rain_filename_pattern": RAIN_FILENAME_PATTERN,
        "transpiration_filename_pattern": TRANSPIRATION_FILENAME_PATTERN,
        "soil_interpolation": SOIL_INTERPOLATION_FILE,
        "soil_source_files": SOIL_SOURCE_FILES,
        "som_source_unit_assumed": SOM_SOURCE_UNIT,
        "som_unit_confirmed": SOM_UNIT_CONFIRMED,
        "initial_tn_mg_L": INITIAL_TN_MG_L,
        "event_rule": EVENT_RULE, "run_label": RUN_LABEL,
        "events": FERTILIZER_EVENTS, "coefficients": COEFFICIENTS,
        "source_files": SOURCE_FILES,
        "nfer_raster": NFER_RASTER,
        "nfer_unit_factor": NFER_UNIT_FACTOR,
        "nfer_to_rice_season_factor": NFER_TO_RICE_SEASON_FACTOR,
        "nfer_mapping_status": NFER_MAPPING_STATUS,
    })


def resolve_path(root, value):
    """绝对路径直接使用；相对路径从项目根目录查找。"""
    return (root / Path(value)).resolve()


def check_config(config, root):
    """在开始运算前指出日期、路径及参数中的具体问题。"""
    start = date.fromisoformat(config['start_date'])
    end = date.fromisoformat(config['end_date'])
    if start.year != 2022 or end.year != 2022 or start > end:
        raise ValueError('START_DATE和END_DATE需在2022年内，且开始日期不晚于结束日期。')
    if config['event_rule'] not in ('latest_event', 'sum_events'):
        raise ValueError('EVENT_RULE请填写latest_event或sum_events。')
    if config['som_source_unit_assumed'] not in ('percent', 'g/kg'):
        raise ValueError('SOM_SOURCE_UNIT请填写percent或g/kg。')
    initial = config['initial_tn_mg_L']
    if not math.isfinite(initial) or initial < 0:
        raise ValueError('INITIAL_TN_MG_L需为大于等于0的有限数值。')
    for name, lower, upper in [('longitude', -180, 180), ('latitude', -90, 90)]:
        if not math.isfinite(config[name]) or not lower <= config[name] <= upper:
            raise ValueError(f'{name}坐标超出有效范围。')
    for key in ('observations', 'soil_interpolation', 'rain_directory', 'transpiration_directory'):
        path = resolve_path(root, config[key])
        if not path.exists():
            raise FileNotFoundError(f'请检查{key}路径：{path}')
    for name in config['source_files']:
        if not resolve_path(root, name).is_file():
            raise FileNotFoundError(f'请检查SOURCE_FILES中的来源文件：{name}')
    if 'nfer_raster' not in config:
        raise ValueError('此配置需提供nfer_raster及各施肥事件的N_fraction。请使用当前脚本生成参数快照。')
    if not resolve_path(root, config['nfer_raster']).is_file():
        raise FileNotFoundError(f'施氮栅格路径不存在：{config["nfer_raster"]}')
    for key in ('nfer_unit_factor', 'nfer_to_rice_season_factor'):
        if not math.isfinite(config[key]) or config[key] <= 0:
            raise ValueError(f'{key}需为大于0的有限数值。')
    for treatment in ('RM', 'RCC'):
        fractions = [e['N_fraction'] for e in config['events'] if e['treatment']==treatment]
        if not fractions or any(not math.isfinite(f) or f <= 0 for f in fractions) or not math.isclose(sum(fractions), 1., abs_tol=1e-9):
            raise ValueError(f'{treatment}各次N_fraction须大于0且合计为1。')
    for event in config['events']:
        event_date = date.fromisoformat(event['date'])
        if event_date.year != 2022:
            raise ValueError(f'施肥日期需在2022年内：{event["date"]}')
        if event['stage'] not in config['coefficients']:
            raise ValueError(f'施肥阶段填写错误：{event["stage"]}')
        if event_date > end:
            print(f'提示：{event["treatment"]} {event["date"]}施肥在结束日期之后，本次不会触发。')
    output_parent = resolve_path(root, config.get('output_directory', '.codex-work/qgis_point_2022/runs'))
    approved_results = (root/'模型/2026.9.3/2022点位验证').resolve()
    if not (output_parent.is_relative_to((root/'.codex-work').resolve()) or output_parent == approved_results):
        raise ValueError('输出目录需为项目.codex-work内路径或指定的2022点位验证结果文件夹。')
    return start, end, output_parent


def read_csv(path):
    """读取带中文表头的CSV，保留源字段。"""
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def write_csv(path, rows):
    if not rows:
        return
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def sha(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(1024*1024), b''):
            h.update(chunk)
    return h.hexdigest()


def sample(path, lon, lat, variable):
    """提取指定坐标的像元，先识别缺失值，再应用缩放系数一次。"""
    from osgeo import gdal, osr
    gdal.UseExceptions()
    ds = gdal.Open(str(path))
    band = ds.GetRasterBand(1)
    src = osr.SpatialReference()
    src.ImportFromEPSG(4326)
    src.SetAxisMappingStrategy(osr.OAMS_TRADITIONAL_GIS_ORDER)
    crs_status = 'embedded_crs'
    if ds.GetProjection():
        dst = osr.SpatialReference()
        dst.ImportFromWkt(ds.GetProjection())
        dst.SetAxisMappingStrategy(osr.OAMS_TRADITIONAL_GIS_ORDER)
        x, y, _ = osr.CoordinateTransformation(src, dst).TransformPoint(lon, lat)
    elif variable == 'T':
        x, y = lon, lat
        crs_status = 'EPSG4326_assumed_from_source_lon_lat_coordinates'
    else:
        raise ValueError('Missing CRS: '+str(path))
    px, py = gdal.ApplyGeoTransform(gdal.InvGeoTransform(ds.GetGeoTransform()), x, y)
    col, row = math.floor(px), math.floor(py)
    if not (0 <= col < ds.RasterXSize and 0 <= row < ds.RasterYSize):
        raise ValueError('Site outside raster: '+str(path))
    raw = float(band.ReadAsArray(col, row, 1, 1)[0, 0])
    nodata = band.GetNoDataValue()
    valid_mask = int(band.GetMaskBand().ReadAsArray(col, row, 1, 1)[0, 0])
    missing = not math.isfinite(raw) or valid_mask == 0 or (nodata is not None and raw == nodata)
    scale, offset = band.GetScale(), band.GetOffset()
    scale = 1. if scale is None else float(scale)
    offset = 0. if offset is None else float(offset)
    meta = dict(ds.GetMetadata())
    meta.update(band.GetMetadata())
    if variable == 'T':
        identity = ' '.join(str(v) for k, v in meta.items() if 'long_name' in k)
        if 'Transpiration' not in identity or not math.isclose(scale, .001, rel_tol=1e-6):
            raise ValueError('Inspect T variable identity and scale: '+str(path))
    value = None if missing else raw*scale+offset
    if value is not None and value < 0:
        raise ValueError('Negative weather value: '+str(path))
    stat = path.stat()
    ds = None
    return dict(variable=variable, file=str(path), size_bytes=stat.st_size, mtime_ns=stat.st_mtime_ns,
                row=row, col=col, raw=raw, scale=scale, offset=offset, nodata=nodata,
                value_mm_day=value, status='missing' if missing else 'valid', crs_status=crs_status)


def prepare_events(config, ph, som):
    """按各次纯氮量、土壤条件和施肥阶段计算C0及响应日期。"""
    events, seen = [], set()
    for item in config['events']:
        e = dict(item)
        key = e['treatment'], e['date']
        if key in seen or e['treatment'] not in ('RM', 'RCC'):
            raise ValueError('Duplicate event or unexpected treatment')
        seen.add(key)
        n, lag = float(e['N_kg_ha']), e['response_lag_days']
        if not math.isfinite(n) or n < 0 or not isinstance(lag, int) or lag < 0:
            raise ValueError('Invalid event N or response lag')
        c = config['coefficients'][e['stage']]
        raw = c['a0']+c['aN']*n+c['apH']*ph+c['aSOM']*som if n > 0 else 0.
        if not math.isfinite(raw) or not math.isfinite(c['k_day']) or c['k_day'] <= 0:
            raise ValueError('Invalid concentration parameters')
        e.update(C0_raw_mg_L=raw, C0_mg_L=max(0., raw),
                 C0_status='clipped_at_zero' if raw < 0 else 'positive', k_day=c['k_day'],
                 pH=ph, SOM_percent_assumed=som,
                 response_date=(date.fromisoformat(e['date'])+timedelta(days=lag)).isoformat())
        events.append(e)
    return sorted(events, key=lambda e: (e['treatment'], e['response_date']))


def predict(day, treatment, events, rule, initial):
    """C=C0 exp[-k(t-1)]；默认施肥后第1天为响应日，t=1。"""
    active = [e for e in events if e['treatment'] == treatment and e['response_date'] <= day.isoformat()]
    if not active:
        return initial, 'initial_state_assumption'
    if rule == 'latest_event':
        active = active[-1:]
    elif rule != 'sum_events':
        raise ValueError('Unknown event_rule')
    value = sum(e['C0_mg_L']*math.exp(-e['k_day']*(day-date.fromisoformat(e['response_date'])).days) for e in active)
    return value, 'event_response'


def assign_raster_nitrogen(config, root):
    """提取施氮像元，进行单位及稻季映射，再分配到各次施肥。"""
    path = resolve_path(root, config['nfer_raster'])
    sampled = sample(path, config['longitude'], config['latitude'], 'Nfer')
    source_value = sampled.pop('value_mm_day')
    if source_value is None:
        raise ValueError('点位施氮栅格为NoData，请核实坐标或提供有依据的补值方案。')
    annual = source_value * config['nfer_unit_factor']
    seasonal = annual * config['nfer_to_rice_season_factor']
    audit = dict(sampled, source_value=source_value, source_unit='g N/m2 land/year',
                 unit_factor=config['nfer_unit_factor'], annual_N_kg_ha_land=annual,
                 rice_season_factor=config['nfer_to_rice_season_factor'],
                 season_N_proxy_kg_ha=seasonal, mapping_status=config['nfer_mapping_status'],
                 sampling='containing_pixel_in_native_grid', raster_sha256=sha(path))
    derived = copy.deepcopy(config)
    for event in derived['events']:
        event['N_kg_ha'] = seasonal * event['N_fraction']
        event['N_status'] = config['nfer_mapping_status']
    return derived, audit


def metrics(obs, pred):
    """计算观测与预测的直接误差；Pearson_r2仅反映相关程度。"""
    n = len(obs)
    if not n:
        return dict(n=0, RMSE_mg_L=None, MAE_mg_L=None, bias_mg_L=None, NSE=None, Pearson_r2=None)
    errors = [p-o for o, p in zip(obs, pred)]
    mo, mp = sum(obs)/n, sum(pred)/n
    sse = sum(e*e for e in errors)
    sso, ssp = sum((o-mo)**2 for o in obs), sum((p-mp)**2 for p in pred)
    cross = sum((o-mo)*(p-mp) for o, p in zip(obs, pred))
    return dict(n=n, RMSE_mg_L=math.sqrt(sse/n), MAE_mg_L=sum(map(abs, errors))/n,
                bias_mg_L=sum(errors)/n, NSE=1-sse/sso if sso > 0 else None,
                Pearson_r2=cross*cross/(sso*ssp) if sso > 0 and ssp > 0 else None)


def pair_observations(observations, predictions, events):
    """按处理和日期配对；记录未入选评价的具体原因。"""
    pairs, seen = [], set()
    event_dates = {(e['treatment'], e['date']) for e in events}
    for row in observations:
        if row['record_id'] in seen:
            raise ValueError('Duplicate observation ID')
        seen.add(row['record_id'])
        if row['role'] != 'observed' or row['metric'] != 'TN' or row['unit'] != 'mg/L':
            raise ValueError('Observation identity or unit mismatch')
        key = row['treatment'], row['date']
        prediction, reasons = predictions.get(key), []
        if row['date_status'] != 'printed_tick' or row['date_exact'] != row['date']:
            reasons.append('date_pending')
        if row['value_status'] not in ('validated_geometry', 'user_revised_workbook'):
            reasons.append('value_pending')
        if key in event_dates:
            reasons.append('fertilizer_sampling_order_pending')
        if prediction is None:
            reasons.append('outside_simulation')
        elif prediction['prediction_status'] == 'initial_state_assumption':
            reasons.append('initial_TN_pending')
        observed = float(row['value'])
        if not math.isfinite(observed) or observed < 0:
            raise ValueError('Invalid observed value')
        simulated = prediction['TN_mg_L'] if prediction else None
        pairs.append(dict(row, simulated_TN_mg_L=simulated,
                          residual_mg_L=simulated-observed if simulated is not None else None,
                          included_in_metrics=not reasons,
                          evaluation_status='included' if not reasons else ';'.join(reasons)))
    return pairs


def make_plot(output, daily, paired):
    # Figure API keeps the QGIS GUI's existing matplotlib backend unchanged.
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    fig = Figure(figsize=(11, 7))
    FigureCanvasAgg(fig)
    axes = fig.subplots(2, 1, sharex=True)
    for ax, treatment in zip(axes, ('RM', 'RCC')):
        rows = [r for r in daily if r['treatment'] == treatment]
        ax.plot([date.fromisoformat(r['date']) for r in rows], [r['TN_mg_L'] for r in rows], label='Literature-parameter prediction', color='#167a9b')
        for included, marker, label in [(True, 'o', 'Scored observations'), (False, 'x', 'Other digitized observations')]:
            points = [r for r in paired if r['treatment'] == treatment and r['included_in_metrics'] == included]
            ax.scatter([date.fromisoformat(r['date']) for r in points], [float(r['value']) for r in points], marker=marker, s=22, label=label)
        ax.set(title=treatment, ylabel='TN (mg/L)')
        ax.grid(alpha=.2)
        ax.legend(fontsize=8)
    fig.suptitle('Qianjiang 2022: raster-driven TN prediction')
    fig.autofmt_xdate()
    fig.tight_layout()
    fig.savefig(output/'TN_comparison.png', dpi=160)


def load_qgis_tables(output):
    from qgis.core import QgsProject, QgsVectorLayer
    from qgis.PyQt.QtCore import QUrl
    for name in ('daily_simulation', 'observation_pairs', 'metrics', 'event_parameters', 'nitrogen_input'):
        uri = QUrl.fromLocalFile(str(output/(name+'.csv'))).toString()+'?type=csv&geomType=none&detectTypes=yes&encoding=UTF-8'
        layer = QgsVectorLayer(uri, name+'_'+output.name, 'delimitedtext')
        if not layer.isValid():
            raise RuntimeError('QGIS table load failed: '+name)
        QgsProject.instance().addMapLayer(layer)


def export_excel_results(output):
    """将同一次运行的CSV汇总为有明确数值/日期/文本类型的Excel工作簿。"""
    from openpyxl import Workbook

    tables = [
        ('metrics', '验证指标', '先看RMSE、MAE和NSE判断数值误差；Pearson相关系数平方只表示相关程度。'),
        ('observation_pairs', '观测对照', '观测来自图件数字化；是否参与评价及筛选原因见included_in_metrics、evaluation_status列。'),
        ('daily_simulation', '逐日模拟', 'RM与RCC逐日田面水TN浓度；单位mg/L。观测时段外属于模型外推。'),
        ('event_parameters', '施肥参数', '栅格施氮代理量按事件比例分配，得到每次N及初始浓度C0。'),
        ('nitrogen_input', '施氮栅格输入', '追溯施氮像元原值、单位换算及稻季映射。年度土地强度到稻季田块口径仍有映射假设。'),
        ('point_weather', '逐日气象', '降雨和蒸腾T用于后续水量计算；T不是完整ET。空白ET表示尚缺输入。'),
        ('raster_samples', '气象读取记录', '气象文件路径、像元位置、缩放系数与原始值，用于复现核查。'),
    ]
    numeric = set('TN_mg_L value simulated_TN_mg_L residual_mg_L n RMSE_mg_L MAE_mg_L bias_mg_L NSE Pearson_r2 N_fraction N_kg_ha C0_mg_L C0_raw_mg_L k_day response_lag_days pH SOM_percent_assumed rainfall_mm_day transpiration_mm_day ET_mm_day source_value annual_N_kg_ha_land season_N_proxy_kg_ha unit_factor rice_season_factor raw scale offset row col value_mm_day nodata pixel_x pixel_y same_day_index sample_size source_row source_value_raw unit_factor digitized_value pixel_value_uncertainty size_bytes'.split())
    dates = {'date', 'date_exact', 'date_lower', 'date_upper', 'date_alternative', 'response_date'}
    booleans = {'included_in_metrics', 'near_zero'}
    workbook = Workbook()
    workbook.remove(workbook.active)
    for filename, title, description in tables:
        path = output/(filename+'.csv')
        rows = read_csv(path)
        if not rows:
            continue
        sheet = workbook.create_sheet(title)
        headers = list(rows[0])
        sheet.append(headers)
        for row_index, record in enumerate(rows, 2):
            for col, key in enumerate(headers, 1):
                raw = record[key]
                cell = sheet.cell(row_index, col)
                if raw == '':
                    continue
                if key in dates:
                    cell.value = datetime.strptime(raw, '%Y-%m-%d')
                    cell.number_format = 'yyyy-mm-dd'
                elif key in booleans:
                    cell.value = raw == 'True'
                elif key in numeric:
                    cell.value = float(raw)
                    if not math.isfinite(cell.value):
                        cell.value = raw
                        cell.data_type = 's'
                    else:
                        cell.number_format = '0.00'
                else:
                    # IDs、校验值、纳秒时间戳和路径明确保存为文本，保留长数字。
                    cell.value = raw
                    cell.data_type = 's'
    destination = output/'2022点位验证结果.xlsx'
    workbook.save(destination)
    print('Excel结果：'+str(destination))
    return destination


def main(config_path=None, load_layers=True):
    """完整运行。默认使用脚本顶部参数；传入config_path可复现指定JSON快照。"""
    # 第1步：确定本次参数和输出位置。
    if config_path is None:
        config = build_config()
        root = resolve_path(HERE, config['project_root'])
        parameter_source = Path(__file__)
    else:
        parameter_source = Path(config_path).resolve()
        config = json.loads(parameter_source.read_text(encoding='utf-8-sig'))
        root = resolve_path(parameter_source.parent, config['project_root'])
    start, end, output_parent = check_config(config, root)
    # 快照使用绝对根目录，移动到运行子目录后仍可定位相同输入。
    config['project_root'] = str(root)
    run_id = datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    public_output = output_parent/run_id
    output = root/'.codex-work/qgis_point_2022/run_records'/run_id
    output.mkdir(parents=True)
    (output/'config_snapshot.json').write_text(json.dumps(config, ensure_ascii=False, indent=2), encoding='utf-8')
    # 第2步：读取土壤输入，核对点位及栅格来源，生成施肥事件。
    soil_path = resolve_path(root, config['soil_interpolation'])
    soil = {r['variable']: r for r in json.loads(soil_path.read_text(encoding='utf-8'))}
    for variable, item in soil.items():
        if abs(item['longitude']-config['longitude']) > 1e-8 or abs(item['latitude']-config['latitude']) > 1e-8:
            raise ValueError('点位坐标与土壤插值文件不一致；修改点位后请更换相应插值文件。')
        source_path = resolve_path(root, config.get('soil_source_files', {}).get(variable, item['source_file']))
        if sha(source_path) != item['source_sha256']:
            raise ValueError(f'{variable}源栅格内容发生变化，请重新生成土壤插值。')
    ph, som_raw = soil['pH']['estimate'], soil['SOM']['estimate']
    som = som_raw*{'percent': 1., 'g/kg': .1}[config['som_source_unit_assumed']]
    derived_config, nitrogen_audit = assign_raster_nitrogen(config, root)
    write_csv(output/'nitrogen_input.csv', [nitrogen_audit])
    events = prepare_events(derived_config, ph, som)
    write_csv(output/'event_parameters.csv', events)
    observations_path = resolve_path(root, config['observations'])
    if observations_path.suffix.lower() == '.xlsx':
        loader = runpy.run_path(str(HERE/'observation_input.py'))
        observations = loader['load_workbook_observations'](
            observations_path, config.get('observation_sheet', '田面水TN浓度变化'))
    else:
        observations = read_csv(observations_path)
    # 保存实际参与本次计算的数据，后续源工作簿修改不会改变本次记录。
    write_csv(output/'observations_input.csv', observations)
    import shutil
    shutil.copy2(observations_path, output/('observation_source'+observations_path.suffix))
    print(f'观测输入：{observations_path}；记录数：{len(observations)}', flush=True)
    calibration = None
    if config.get('calibrate_k', False):
        calibration_module = runpy.run_path(str(HERE/'bounded_k.py'))
        calibration = calibration_module['calibrate'](events, observations, config['k_bounds'])
        (output/'k_calibration.json').write_text(json.dumps(calibration, ensure_ascii=False, indent=2), encoding='utf-8')
        write_csv(output/'event_parameters.csv', events)
        print('k率定结果：'+json.dumps(calibration['parameters'], ensure_ascii=False), flush=True)
    ledger, weather, daily = [], [], []
    count = (end-start).days+1
    # 第3步：逐日提取气象；按处理计算TN浓度。
    for i in range(count):
        day = start+timedelta(days=i)
        doy = day.timetuple().tm_yday
        met = {}
        rain_pattern = config.get('rain_filename_pattern', 'PRE_{year}_Day_{doy:03d}.tif')
        transpiration_pattern = config.get('transpiration_filename_pattern', 'T_{year}{doy:03d}.tif')
        for var, folder, pattern in [('P', 'rain_directory', rain_pattern), ('T', 'transpiration_directory', transpiration_pattern)]:
            name = pattern.format(year=day.year, doy=doy, date=day.isoformat())
            item = sample(resolve_path(root, config[folder])/name, config['longitude'], config['latitude'], var)
            ledger.append(dict(date=day.isoformat(), **item))
            met[var] = item['value_mm_day']
        weather.append(dict(date=day.isoformat(), rainfall_mm_day=met['P'], transpiration_mm_day=met['T'], ET_mm_day=None,
                            role='water_balance_input; not_used_in_TN_equation'))
        for treatment in ('RM', 'RCC'):
            value, status = predict(day, treatment, events, config['event_rule'], config['initial_tn_mg_L'])
            daily.append(dict(date=day.isoformat(), treatment=treatment, TN_mg_L=value, prediction_status=status,
                              event_rule=config['event_rule'], run_label=config['run_label']))
        if i % 30 == 0 or i == count-1:
            print(f'[{i+1}/{count}] {day}', flush=True)
    # 第4步：与观测配对，同处理同日均值计为一个评价样本。
    predictions = {(r['treatment'], r['date']): r for r in daily}
    paired = pair_observations(observations, predictions, events)
    training_keys = set(tuple(key) for key in calibration['training_keys']) if calibration else set()
    for row in paired:
        row['calibration_role'] = ('用于k率定' if (row['treatment'],row['date']) in training_keys
                                   else '未参与k率定') if calibration else '文献参数预测'
    scores = []
    for treatment in ('RM', 'RCC'):
        selected = [r for r in paired if r['treatment'] == treatment and r['included_in_metrics']]
        dates = sorted({r['date'] for r in selected})
        obs = [sum(float(r['value']) for r in selected if r['date'] == d)/sum(r['date'] == d for r in selected) for d in dates]
        pred = [predictions[treatment, d]['TN_mg_L'] for d in dates]
        scores.append(dict(treatment=treatment, unit='mg/L', aggregation='one_mean_per_sampling_date', run_label=config['run_label'], **metrics(obs, pred)))
    # 第5步：导出表格、图件、参数和文件校验值。
    for name, rows in [('daily_simulation', daily), ('point_weather', weather), ('raster_samples', ledger), ('observation_pairs', paired), ('metrics', scores)]:
        write_csv(output/(name+'.csv'), rows)
    if calibration:
        calibration_module['write_comparison'](output, paired, events, calibration, metrics, write_csv)
    # 简洁成果与完整记录共用相同计算数据。
    report_script = HERE/'result_report.py'
    report = runpy.run_path(str(report_script), run_name='point_result_report')
    report['publish_results'](output, public_output)
    sources = [parameter_source, Path(__file__), soil_path, observations_path]+[resolve_path(root, p) for p in config['source_files']]
    sources.append(report_script)
    sources.append(HERE/'observation_input.py')
    sources.append(HERE/'bounded_k.py')
    nfer_path = resolve_path(root, config['nfer_raster'])
    sources.extend([nfer_path] + [p for p in (Path(str(nfer_path)+'.hdr'), Path(str(nfer_path)+'.aux.xml')) if p.exists()])
    manifest = dict(run_label=config['run_label'], scope='TN concentration; bounded k calibration' if calibration else 'TN concentration; literature-parameter trial', python=sys.version,
                    public_output=str(public_output), archive_directory=str(output),
                    n_days=count, n_predictions=len(daily), scored_dates={r['treatment']: r['n'] for r in scores},
                    missing_weather_samples=sum(r['status']=='missing' for r in ledger),
                    source_sha256={str(p): sha(p) for p in sources}, output_sha256={p.name: sha(p) for p in output.glob('*.csv')},
                    raster_fingerprint='path, size, mtime, pixel, scale and value; not full-file hashes',
                    evidence_status='2018 annual land N mapped to 2022 rice-season proxy; crop-area mapping, SOM unit and RM second-event stage pending',
                    calibration_scope='first-event day 1-7; fixed raster C0; main metrics include training dates' if calibration else 'none')
    (output/'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    manifest['public_output_sha256'] = {p.name:sha(p) for p in public_output.iterdir() if p.is_file()}
    (output/'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    if load_layers:
        load_qgis_tables(output)
    print('\n查看结果：'+str(public_output))
    print('完整复现记录：'+str(output))
    print(json.dumps(scores, ensure_ascii=False, indent=2))
    print('请先看Excel的验证指标和误差诊断；参数口径见复现位置。')
    return public_output


# QGIS编辑器通过InteractiveInterpreter执行，环境名为__console__。
# runpy和命令行执行使用__main__；普通导入保留函数供调用。
if __name__ in ('__main__', '__console__'):
    print('开始运行2022年栅格驱动点位计算……', flush=True)
    main(load_layers=LOAD_TABLES_IN_QGIS and '--headless' not in sys.argv)
